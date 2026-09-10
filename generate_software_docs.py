import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

DOCX_PATH = "SWMS_Software_Documentation.docx"
PDF_PATH = "SWMS_Software_Documentation.pdf"

print("Starting document generation...")

# ==========================================
# 1. HELPER FUNCTIONS FOR DOCX FORMATTING
# ==========================================

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = 'Arial'
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(11, 17, 30) # Dark Industrial Blue
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(2, 132, 199) # Sky Blue
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(22, 163, 74) # Emerald Green
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout_box(doc, title, text, border_color_hex="0284c7", bg_color_hex="f0f9ff"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[NOTE] {title}\n")
    r_title.bold = True
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(11, 17, 30)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_docx():
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title Header Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("AI-POWERED SMART WASTE SEGREGATION SYSTEM")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(11, 17, 30)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("SOFTWARE ARCHITECTURE, DIGITAL TWIN DASHBOARD & TELEMETRY REPORT")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(2, 132, 199)
    p_sub.paragraph_format.space_after = Pt(12)

    # Metadata Card Table
    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("System Type:", " Industrial IoT Digital Twin & Safety Control"), ("Frontend Stack:", " React 19 + Vite 8 + Tailwind CSS")],
        [("Cloud IoT Platform:", " ThingSpeak REST API (Channel #3470506)"), ("Telemetry Rate:", " 10-Second Continuous Heartbeat Poller")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (label, val) in enumerate(row):
            cell = meta_tbl.cell(r_idx, c_idx)
            set_cell_background(cell, "f8fafc")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            r1 = p.add_run(label)
            r1.bold = True
            r1.font.name = 'Arial'
            r1.font.size = Pt(9.5)
            r2 = p.add_run(val)
            r2.font.name = 'Arial'
            r2.font.size = Pt(9.5)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY & SOFTWARE SCOPE
    # ----------------------------------------------------
    add_styled_heading(doc, "1. Executive Summary & Software Scope", level=1)
    
    p = doc.add_paragraph()
    p.add_run(
        "This software documentation presents the complete digital architecture, frontend framework design, real-time telemetry ingestion pipelines, safety state evaluation engine, and analytics visualization platform for the "
    )
    r_bold = p.add_run("AI-Powered Smart Waste Segregation Bin — Digital Twin System")
    r_bold.bold = True
    p.add_run(
        ". Built as a production-grade Industrial IoT (IIoT) control center, the software functions as a high-fidelity Digital Twin that renders live fluid capacity levels across multiple physical bin compartments, monitors critical safety parameters (flame, combustible gas/smoke, and ferrous metal proximity), and automatically alerts system administrators when hazardous thresholds are exceeded."
    )
    p.paragraph_format.space_after = Pt(8)

    add_callout_box(
        doc,
        "Edge AI vs. Digital Twin Telemetry Isolation Guarantee",
        "The software architecture enforces strict decoupling between physical hardware intelligence and web visualization. Optical image classification (PET/Plastic, Dry/Recyclable, Organic) is executed 100% locally on the ESP32-CAM microcontroller edge hardware. The React Digital Twin Dashboard does NOT run synthetic inferences or artificial predictions; it serves exclusively as an authoritative, high-fidelity monitoring and safety control station consuming live cloud telemetry stream."
    )

    # ----------------------------------------------------
    # SECTION 2: SYSTEM ARCHITECTURE & DATA FLOW
    # ----------------------------------------------------
    add_styled_heading(doc, "2. System Architecture & Data Flow", level=1)
    
    p = doc.add_paragraph()
    p.add_run("The complete end-to-end data pipeline spans four primary operational layers: Edge Microcontroller Layer, Cloud IoT Data Gateway, React Application Service Layer, and Neumorphic Digital Twin UI Layer.")
    p.paragraph_format.space_after = Pt(8)

    arch_code = """
+-----------------------------------------------------------------------------------+
|                                 SMART WASTE BIN HARDWARE                          |
|  +---------------------------+                +---------------------------------+  |
|  |     ESP32-CAM (Edge AI)   |                | Physical Sensors (HC-SR04, MQ2, |  |
|  | Local Vision Classification|                | Flame Detector, Proximity)      |  |
|  +-------------+-------------+                +----------------+----------------+  |
+----------------|-----------------------------------------------+------------------+
                 |                                               |
                 +-----------------------+-----------------------+
                                         |
                                         v
                            +--------------------------+
                            | ESP32 Main IoT Controller|
                            +------------+-------------+
                                         | HTTP POST (JSON / Form)
                                         v
                            +--------------------------+
                            |   ThingSpeak Cloud IoT   |
                            |   (REST Field Channels)  |
                            +------------+-------------+
                                         | HTTP GET (10s Polling Hook)
                                         v
                            +--------------------------+
                            | React 19 Digital Twin    |
                            | Dashboard (Vite Frontend)|
                            +--------------------------+
"""
    p_code = doc.add_paragraph()
    r_code = p_code.add_run(arch_code)
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(30, 41, 59)
    p_code.paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 3: TECH STACK & DEPENDENCY MATRIX
    # ----------------------------------------------------
    add_styled_heading(doc, "3. Software Technology Stack & Library Inventory", level=1)
    
    p = doc.add_paragraph()
    p.add_run("The application is constructed using a modern React 19 ecosystem bundled with Vite 8. Below is the comprehensive software dependency matrix:")
    p.paragraph_format.space_after = Pt(8)

    stack_tbl = doc.add_table(rows=1, cols=4)
    stack_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Package Name", "Version", "Role / Purpose", "Architectural Scope"]
    hdr_cells = stack_tbl.rows[0].cells
    for i, title in enumerate(headers):
        set_cell_background(hdr_cells[i], "0b111e")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    stack_data = [
        ["React", "^19.2.8", "Core Component Engine & Declarative UI State", "Frontend Framework"],
        ["Vite", "^8.2.2", "Lightning-fast HMR Build Tool & Module Bundler", "Build System"],
        ["Tailwind CSS", "^3.4.17", "Utility-First CSS with Dark Neumorphic Tokens", "Styling Engine"],
        ["Axios", "^1.20.0", "Promise-based HTTP Client for ThingSpeak REST API", "Data Ingestion"],
        ["Recharts", "^3.10.1", "SVG Time-Series Analytics & Telemetry Line/Area Charts", "Data Visualization"],
        ["Lucide React", "^1.38.0", "Industrial IoT SVG Vector Icon Set", "User Interface"],
        ["Framer Motion", "^13.1.1", "Smooth Micro-animations & Fluid Container Transitions", "UI Animation"],
        ["Oxlint", "^1.79.0", "High-performance JS/JSX Code Quality & Linter", "Code Maintenance"],
    ]

    for row_idx, row in enumerate(stack_data):
        r_cells = stack_tbl.add_row().cells
        for i, val in enumerate(row):
            set_cell_background(r_cells[i], "f8fafc" if row_idx % 2 == 0 else "ffffff")
            set_cell_margins(r_cells[i], top=80, bottom=80, left=100, right=100)
            p = r_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # SECTION 4: CLOUD IOT COMMUNICATION & THINGSPEAK SCHEMA
    # ----------------------------------------------------
    add_styled_heading(doc, "4. Cloud IoT Communication & ThingSpeak Schema", level=1)
    
    p = doc.add_paragraph()
    p.add_run(
        "Communication between physical bin sensors and the software dashboard is mediated through the MathWorks ThingSpeak IoT Cloud platform (Channel ID: "
    )
    p.add_run("3470506").bold = True
    p.add_run(
        "). The frontend queries the REST channels asynchronously every 10 seconds to retrieve incoming sensor telemetry feeds."
    )
    p.paragraph_format.space_after = Pt(8)

    # Schema Table
    schema_tbl = doc.add_table(rows=1, cols=5)
    schema_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    s_headers = ["ThingSpeak Field", "Physical Sensor Name", "Target Metric", "Valid Range / Format", "UI Destination"]
    s_hdr_cells = schema_tbl.rows[0].cells
    for i, title in enumerate(s_headers):
        set_cell_background(s_hdr_cells[i], "0284c7")
        set_cell_margins(s_hdr_cells[i], top=100, bottom=100, left=100, right=100)
        p = s_hdr_cells[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    schema_data = [
        ["Field 1", "Ultrasonic Sensor 1", "Bin 01 Fill Level", "0% - 100% (or distance cm)", "BIN 01 — PET / PLASTIC"],
        ["Field 2", "Ultrasonic Sensor 2", "Bin 02 Fill Level", "0% - 100% (or distance cm)", "BIN 02 — DRY / RECYCLABLE"],
        ["Field 3", "Ultrasonic Sensor 3", "Bin 03 Fill Level", "0% - 100% (or distance cm)", "BIN 03 — ORGANIC / GENERAL"],
        ["Field 4", "Optical Flame Sensor", "Fire / Spark Hazard", "0 = Safe, 1 = Flame Alert", "Flame Sensor & Emergency Banner"],
        ["Field 5", "Gas / Smoke (MQ-2)", "Air Quality / Smoke PPM", "0 - 1000+ PPM", "Radial Gas Gauge (<300 Safe, >=800 Danger)"],
        ["Field 6", "Inductive Proximity", "Ferrous Metal Detection", "0 = No Metal, 1 = Metal", "Proximity Metal Indicator Card"],
    ]

    for row_idx, row in enumerate(schema_data):
        r_cells = schema_tbl.add_row().cells
        for i, val in enumerate(row):
            set_cell_background(r_cells[i], "f0f9ff" if row_idx % 2 == 0 else "ffffff")
            set_cell_margins(r_cells[i], top=80, bottom=80, left=100, right=100)
            p = r_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # REST Endpoints Specification
    add_styled_heading(doc, "ThingSpeak REST API Integration Endpoints", level=2)
    
    ep_p = doc.add_paragraph()
    ep_p.add_run("1. Fetch Latest Telemetry Entry:\n").bold = True
    ep_p.add_run("   GET https://api.thingspeak.com/channels/3470506/feeds/last.json?api_key=W2K6VI4WUKVWX3N3\n").font.name = 'Consolas'
    ep_p.add_run("2. Fetch Historical Feeds (40 frames for time-series charts):\n").bold = True
    ep_p.add_run("   GET https://api.thingspeak.com/channels/3470506/feeds.json?api_key=W2K6VI4WUKVWX3N3&results=40\n").font.name = 'Consolas'
    ep_p.paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 5: SOFTWARE LOGIC & ALGORITHMS
    # ----------------------------------------------------
    add_styled_heading(doc, "5. Software Logic, Algorithms & Mathematical Models", level=1)
    
    p = doc.add_paragraph()
    p.add_run("The software framework encapsulates several dedicated algorithmic models within `sensorUtils.js` and `alertUtils.js` to process raw telemetry values into actionable operational states.")
    p.paragraph_format.space_after = Pt(8)

    # Algorithm 1: Distance Calibration
    add_styled_heading(doc, "Algorithm 1: Ultrasonic Fill Level Calibration Formula", level=2)
    p = doc.add_paragraph()
    p.add_run("When operating under RAW_DISTANCE calibration mode, the fill percentage is calculated using linear inverse transformation bounded between empty (40cm) and full (5cm) thresholds:")
    p.paragraph_format.space_after = Pt(4)

    math_p = doc.add_paragraph()
    r_math = math_p.add_run("   Fill % = min(100, max(0, ((Empty_Distance - Current_Distance) / (Empty_Distance - Full_Distance)) * 100))")
    r_math.font.name = 'Consolas'
    r_math.font.size = Pt(9.5)
    r_math.font.bold = True
    r_math.font.color.rgb = RGBColor(2, 132, 199)
    math_p.paragraph_format.space_after = Pt(8)

    # Algorithm 2: Tri-State Safety System
    add_styled_heading(doc, "Algorithm 2: Tri-State Global Safety System Decision Engine", level=2)
    p = doc.add_paragraph()
    p.add_run("The system evaluates environmental safety parameters simultaneously using a hierarchical priority rule matrix:")
    p.paragraph_format.space_after = Pt(4)

    safety_rules = [
        ("EMERGENCY STATE (Priority 1):", " Triggered if Flame Sensor == 1 (Field 4) OR Gas Concentration >= 800 PPM (Field 5). Result: Halts sorting, plays animated emergency alert overlay."),
        ("WARNING STATE (Priority 2):", " Triggered if Any Bin Fill Level >= 80% OR Gas Concentration is between 300 PPM and 799 PPM. Result: Displays warning badge and prepares collection dispatch."),
        ("NORMAL STATE (Priority 3):", " Triggered when Flame == 0 AND Gas < 300 PPM AND All Bins < 80%. Result: Nominal continuous monitoring.")
    ]
    for title, desc in safety_rules:
        bp = doc.add_paragraph()
        bp.paragraph_format.left_indent = Inches(0.2)
        bp.paragraph_format.space_after = Pt(3)
        r1 = bp.add_run(f"• {title}")
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)
        r2 = bp.add_run(desc)
        r2.font.name = 'Arial'
        r2.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Algorithm 3: Non-Duplicative Event Logger
    add_styled_heading(doc, "Algorithm 3: Non-Duplicative State Transition Logger Engine", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "To prevent redundant notification floods during continuous 10-second polling, `alertUtils.js` tracks previous state hashes. A new event entry is logged ONLY when a genuine threshold boundary is crossed (e.g., transition from SAFE to HIGH smoke, or NORMAL to ALMOST FULL bin capacity). Event logs are stored in a ring-buffer capped at 50 historical entries."
    )
    p.paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 6: USER INTERFACE COMPONENTS & DIGITAL TWIN
    # ----------------------------------------------------
    add_styled_heading(doc, "6. UI Component Architecture & Digital Twin Visualizers", level=1)
    
    p = doc.add_paragraph()
    p.add_run("The user interface follows a modular, industrial control center design pattern using Neumorphic and Glassmorphic aesthetics. Key components include:")
    p.paragraph_format.space_after = Pt(8)

    comp_list = [
        ("Header.jsx", "Top navigation bar displaying live pulse indicator, last sync timestamp, channel ID badge, and manual telemetry sync button."),
        ("SystemStatus.jsx", "Global banner rendering System State (NORMAL, WARNING, EMERGENCY) and high-priority fire/gas warning alerts."),
        ("KPIGrid.jsx", "Top metrics grid featuring 7 cards: System State, Active Bins, Highest Fill %, Fire Alert, Gas Concentration, Metal Detection, and Sensor Signal Health."),
        ("DigitalTwinBin.jsx", "Master 3D physical bin cylinder visualization with dynamic fluid wave meniscus, height percentages, HC-SR04 bracket graphic, and fill markings (25%, 50%, 75%, 80%, 95%, 100%)."),
        ("SafetyPanel.jsx & Components", "Dedicated multi-sensor safety card container encapsulating FlameSensor.jsx, SmokeGauge.jsx (animated radial SVG gas meter), and MetalDetection.jsx."),
        ("SensorHealth.jsx", "Diagnostic telemetry matrix assessing active vs. inactive fields across all physical sensors."),
        ("CollectionStatus.jsx", "Waste logistics planner providing automated collection recommendations and priority dispatch statuses."),
        ("AlertsPanel.jsx", "Event log timeline with severity filter tabs (All, Critical, Warning, Info)."),
        ("FillLevelChart.jsx & SmokeChart.jsx", "Recharts historical telemetry visualization rendering fill % and smoke PPM time series across 40 frames.")
    ]

    for name, desc in comp_list:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        r_name = p.add_run(f"▪ {name}: ")
        r_name.bold = True
        r_name.font.name = 'Arial'
        r_name.font.size = Pt(10)
        r_name.font.color.rgb = RGBColor(2, 132, 199)
        r_desc = p.add_run(desc)
        r_desc.font.name = 'Arial'
        r_desc.font.size = Pt(10)
        r_desc.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 7: SOURCE CODE DIRECTORY HIERARCHY
    # ----------------------------------------------------
    add_styled_heading(doc, "7. Source Code Directory Structure", level=1)
    
    tree_text = """smart-waste/
├── .env                         # ThingSpeak Channel ID & Read Key configuration
├── index.html                   # HTML entrypoint with Inter/JetBrains fonts
├── package.json                 # Frontend dependencies and Vite build scripts
├── postcss.config.js            # PostCSS configuration
├── tailwind.config.js           # Custom dark industrial neumorphic theme tokens
└── src/
    ├── config/
    │   └── sensorConfig.js      # Sensor thresholds, mode & bin metadata
    ├── services/
    │   └── thingspeak.js        # Axios REST client & feed normalization
    ├── utils/
    │   ├── sensorUtils.js       # Fill calculation, status & parsing utilities
    │   └── alertUtils.js        # Non-duplicative state transition logger
    ├── hooks/
    │   └── useThingSpeak.js     # 10s polling custom hook with window focus pause
    ├── components/
    │   ├── Header.jsx           # Live pulse indicator, sync button, channel ID
    │   ├── SystemStatus.jsx     # Emergency banner & global status bar
    │   ├── KPIGrid.jsx          # 7 key operational metric summary cards
    │   ├── StatusCard.jsx       # Reusable glassmorphic container card
    │   ├── DigitalTwin.jsx      # Master 3D-styled bin container wrapper
    │   ├── DigitalTwinBin.jsx   # Animated fluid cylinder with level markings
    │   ├── SafetyPanel.jsx      # Multi-sensor safety container
    │   ├── FlameSensor.jsx      # Optical fire hazard indicator
    │   ├── SmokeGauge.jsx       # Animated radial SVG gas gauge (0-1000 PPM)
    │   ├── MetalDetection.jsx   # Inductive proximity sensor state card
    │   ├── SensorHealth.jsx     # Telemetry diagnostics grid for all 6 fields
    │   ├── CollectionStatus.jsx # Automated logistics recommendation planner
    │   ├── AlertsPanel.jsx      # Real-time event log with severity filtering
    │   ├── FillLevelChart.jsx   # Recharts fill % historical time series
    │   ├── SmokeChart.jsx       # Recharts gas/smoke historical time series
    │   ├── LoadCellPlaceholder.jsx # Future hardware weight expansion card (-- g)
    │   ├── PowerPlaceholder.jsx    # Future power telemetry card (-- W, -- V)
    │   └── EdgeAINote.jsx       # ESP32-CAM Edge AI architecture diagram
    ├── pages/
    │   └── Dashboard.jsx        # Complete assembled dashboard layout
    ├── App.jsx                  # Root React application component
    ├── main.jsx                 # Vite application entrypoint
    └── index.css                # Tailwind directives & industrial dark styles"""

    p_tree = doc.add_paragraph()
    r_tree = p_tree.add_run(tree_text)
    r_tree.font.name = 'Consolas'
    r_tree.font.size = Pt(8.5)
    r_tree.font.color.rgb = RGBColor(30, 41, 59)
    p_tree.paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # SECTION 8: SETUP, BUILD & DEPLOYMENT MANUAL
    # ----------------------------------------------------
    add_styled_heading(doc, "8. Setup, Build & Deployment Manual", level=1)
    
    p = doc.add_paragraph()
    p.add_run("Follow these commands to configure, build, and deploy the software dashboard:")
    p.paragraph_format.space_after = Pt(6)

    steps = [
        ("Step 1: Environment Setup", "Create a .env file in the project root:\n   VITE_THINGSPEAK_CHANNEL_ID=3470506\n   VITE_THINGSPEAK_READ_API_KEY=W2K6VI4WUKVWX3N3"),
        ("Step 2: Install Node Dependencies", "Execute in shell: npm install"),
        ("Step 3: Launch Local Development Server", "Execute in shell: npm run dev (Launches Vite dev server on http://localhost:5173)"),
        ("Step 4: Execute Production Build", "Execute in shell: npm run build (Outputs optimized static files to /dist)"),
        ("Step 5: Preview Production Bundle", "Execute in shell: npm run preview")
    ]

    for title, cmd in steps:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"• {title}:\n")
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)
        r2 = p.add_run(f"   {cmd}")
        r2.font.name = 'Consolas'
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.save(DOCX_PATH)
    print(f"Generated DOCX successfully: {DOCX_PATH}")

# ==========================================
# 2. HELPER FUNCTIONS FOR REPORTLAB PDF
# ==========================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawString(54, 750, "AI-Powered Smart Waste Segregation Bin -- Software Report")
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)
        
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 36, page_text)
        self.drawString(54, 36, "CONFIDENTIAL -- FOR ACADEMIC & PROJECT REPORT PURPOSES")
        self.line(54, 48, 558, 48)
        self.restoreState()

def build_pdf():
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=64,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0b111e"),
        alignment=1,
        spaceAfter=4
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#0284c7"),
        alignment=1,
        spaceAfter=12
    )

    h1_style = ParagraphStyle(
        'CustomH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#0b111e"),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'CustomH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#0284c7"),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6
    )

    callout_style = ParagraphStyle(
        'CalloutText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#1e293b")
    )

    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#1e293b")
    )

    story = []

    story.append(Paragraph("AI-POWERED SMART WASTE SEGREGATION SYSTEM", title_style))
    story.append(Paragraph("SOFTWARE ARCHITECTURE, DIGITAL TWIN DASHBOARD & TELEMETRY REPORT", sub_style))

    meta_data = [
        [Paragraph("<b>System Type:</b> Industrial IoT Digital Twin", body_style), Paragraph("<b>Frontend Stack:</b> React 19 + Vite 8 + Tailwind CSS", body_style)],
        [Paragraph("<b>Cloud IoT:</b> ThingSpeak REST API (#3470506)", body_style), Paragraph("<b>Telemetry Rate:</b> 10s Continuous Heartbeat Poller", body_style)]
    ]
    t_meta = Table(meta_data, colWidths=[250, 254])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 8))

    # Section 1
    story.append(Paragraph("1. Executive Summary & Software Scope", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    story.append(Paragraph("This software documentation presents the complete digital architecture, frontend framework design, real-time telemetry ingestion pipelines, safety state evaluation engine, and analytics visualization platform for the <b>AI-Powered Smart Waste Segregation Bin -- Digital Twin System</b>. Built as a production-quality Industrial IoT (IIoT) control center, the software functions as a high-fidelity Digital Twin that renders live fluid capacity levels across multiple physical bin compartments, monitors critical safety parameters (flame, combustible gas/smoke, and ferrous metal proximity), and automatically alerts system administrators when hazardous thresholds are exceeded.", body_style))

    callout_data = [[
        Paragraph("<b>[NOTE] Edge AI vs. Digital Twin Telemetry Isolation Guarantee</b><br/><br/>The software architecture enforces strict decoupling between physical hardware intelligence and web visualization. Optical image classification (PET/Plastic, Dry/Recyclable, Organic) is executed 100% locally on the ESP32-CAM microcontroller edge hardware. The React Digital Twin Dashboard does NOT run synthetic inferences or artificial predictions; it serves exclusively as an authoritative, high-fidelity monitoring and safety control station consuming live cloud telemetry stream.", callout_style)
    ]]
    t_callout = Table(callout_data, colWidths=[504])
    t_callout.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f0f9ff")),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
        ('RIGHTPADDING', (0,0), (-1,-1), 10),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LINELEFT', (0,0), (0,0), 3, colors.HexColor("#0284c7")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#bae6fd")),
    ]))
    story.append(t_callout)
    story.append(Spacer(1, 8))

    # Section 2
    story.append(Paragraph("2. System Architecture & Data Flow", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    story.append(Paragraph("The software pipeline integrates four operational layers: Edge Hardware Sensors, ESP32 Microcontroller, ThingSpeak IoT Gateway, and the React Digital Twin Dashboard.", body_style))

    arch_text = """+-----------------------------------------------------------------------------------+
|                                 SMART WASTE BIN HARDWARE                          |
|  +---------------------------+                +---------------------------------+  |
|  |     ESP32-CAM (Edge AI)   |                | Physical Sensors (HC-SR04, MQ2, |  |
|  | Local Vision Classification|                | Flame Detector, Proximity)      |  |
|  +-------------+-------------+                +----------------+----------------+  |
+----------------|-----------------------------------------------+------------------+
                 |                                               |
                 +-----------------------+-----------------------+
                                         |
                                         v
                            +--------------------------+
                            | ESP32 Main IoT Controller|
                            +------------+-------------+
                                         | HTTP POST (JSON / Form)
                                         v
                            +--------------------------+
                            |   ThingSpeak Cloud IoT   |
                            |   (REST Field Channels)  |
                            +------------+-------------+
                                         | HTTP GET (10s Polling Hook)
                                         v
                            +--------------------------+
                            | React 19 Digital Twin    |
                            | Dashboard (Vite Frontend)|
                            +--------------------------+"""
    t_arch = Table([[Paragraph(arch_text.replace("\n", "<br/>").replace(" ", "&nbsp;"), code_style)]], colWidths=[504])
    t_arch.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_arch)
    story.append(Spacer(1, 8))

    # Section 3
    story.append(Paragraph("3. Software Technology Stack & Dependencies", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    
    stack_data = [
        [Paragraph("<font color='white'><b>Package</b></font>", body_style), Paragraph("<font color='white'><b>Version</b></font>", body_style), Paragraph("<font color='white'><b>Purpose / Role</b></font>", body_style), Paragraph("<font color='white'><b>Scope</b></font>", body_style)],
        [Paragraph("React", body_style), Paragraph("^19.2.8", body_style), Paragraph("Core Component Engine & UI State", body_style), Paragraph("Frontend Framework", body_style)],
        [Paragraph("Vite", body_style), Paragraph("^8.2.2", body_style), Paragraph("Module Bundler & Dev Server", body_style), Paragraph("Build Tool", body_style)],
        [Paragraph("Tailwind CSS", body_style), Paragraph("^3.4.17", body_style), Paragraph("Industrial Neumorphic UI Styling", body_style), Paragraph("Styling Engine", body_style)],
        [Paragraph("Axios", body_style), Paragraph("^1.20.0", body_style), Paragraph("Promise-based HTTP REST Client", body_style), Paragraph("Data Ingestion", body_style)],
        [Paragraph("Recharts", body_style), Paragraph("^3.10.1", body_style), Paragraph("SVG Telemetry Time-Series Charts", body_style), Paragraph("Analytics", body_style)],
        [Paragraph("Lucide React", body_style), Paragraph("^1.38.0", body_style), Paragraph("Industrial Vector Icon Library", body_style), Paragraph("UI Icons", body_style)],
        [Paragraph("Framer Motion", body_style), Paragraph("^13.1.1", body_style), Paragraph("Fluid Bin Wave Animations", body_style), Paragraph("Animations", body_style)],
    ]
    t_stack = Table(stack_data, colWidths=[90, 60, 234, 120])
    t_stack.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0b111e")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_stack)
    story.append(Spacer(1, 8))

    # Section 4
    story.append(Paragraph("4. Cloud IoT Communication & ThingSpeak Field Schema", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    
    schema_pdf_data = [
        [Paragraph("<font color='white'><b>Field</b></font>", body_style), Paragraph("<font color='white'><b>Sensor Name</b></font>", body_style), Paragraph("<font color='white'><b>Monitored Metric</b></font>", body_style), Paragraph("<font color='white'><b>Valid Range</b></font>", body_style), Paragraph("<font color='white'><b>UI Destination</b></font>", body_style)],
        [Paragraph("Field 1", body_style), Paragraph("Ultrasonic 1", body_style), Paragraph("Bin 01 Fill Level", body_style), Paragraph("0% - 100%", body_style), Paragraph("BIN 01 — PLASTIC", body_style)],
        [Paragraph("Field 2", body_style), Paragraph("Ultrasonic 2", body_style), Paragraph("Bin 02 Fill Level", body_style), Paragraph("0% - 100%", body_style), Paragraph("BIN 02 — RECYCLABLE", body_style)],
        [Paragraph("Field 3", body_style), Paragraph("Ultrasonic 3", body_style), Paragraph("Bin 03 Fill Level", body_style), Paragraph("0% - 100%", body_style), Paragraph("BIN 03 — ORGANIC", body_style)],
        [Paragraph("Field 4", body_style), Paragraph("Flame Sensor", body_style), Paragraph("Fire / Spark Hazard", body_style), Paragraph("0 = Safe, 1 = Alert", body_style), Paragraph("Emergency Banner", body_style)],
        [Paragraph("Field 5", body_style), Paragraph("Gas (MQ-2)", body_style), Paragraph("Air Quality / Smoke", body_style), Paragraph("0 - 1000+ PPM", body_style), Paragraph("Radial Gas Gauge", body_style)],
        [Paragraph("Field 6", body_style), Paragraph("Proximity", body_style), Paragraph("Ferrous Metal Detection", body_style), Paragraph("0 = No, 1 = Metal", body_style), Paragraph("Metal Indicator Card", body_style)],
    ]
    t_schema = Table(schema_pdf_data, colWidths=[54, 90, 120, 100, 140])
    t_schema.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0284c7")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_schema)
    story.append(Spacer(1, 8))

    # Section 5
    story.append(Paragraph("5. Software Logic, Algorithms & Mathematical Models", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    
    story.append(Paragraph("<b>Algorithm 1: Distance Calibration Equation</b>", h2_style))
    story.append(Paragraph("Fill % = min(100, max(0, ((Empty_Dist - Current_Dist) / (Empty_Dist - Full_Dist)) * 100))", code_style))
    story.append(Spacer(1, 4))

    story.append(Paragraph("<b>Algorithm 2: Tri-State Global Safety System Rules</b>", h2_style))
    story.append(Paragraph("• <b>EMERGENCY STATE:</b> Triggered if Flame == 1 OR Smoke >= 800 PPM. Halts bin sorting.", body_style))
    story.append(Paragraph("• <b>WARNING STATE:</b> Triggered if Any Bin Fill >= 80% OR Smoke 300-799 PPM.", body_style))
    story.append(Paragraph("• <b>NORMAL STATE:</b> Triggered when Flame == 0 AND Smoke < 300 PPM AND All Bins < 80%.", body_style))
    story.append(Spacer(1, 8))

    # Section 6 & 7
    story.append(Paragraph("6. UI Components & Directory Hierarchy", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    story.append(Paragraph("The UI comprises modular glassmorphic cards: <code>Header.jsx</code>, <code>SystemStatus.jsx</code>, <code>KPIGrid.jsx</code>, <code>DigitalTwinBin.jsx</code>, <code>SafetyPanel.jsx</code>, <code>SensorHealth.jsx</code>, <code>CollectionStatus.jsx</code>, <code>AlertsPanel.jsx</code>, and Recharts graph modules.", body_style))
    story.append(Spacer(1, 8))

    story.append(Paragraph("7. Quick Start & Execution Commands", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#0284c7"), spaceAfter=6))
    story.append(Paragraph("1. Configure <code>.env</code> with ThingSpeak Channel ID (3470506) and Read API Key.", body_style))
    story.append(Paragraph("2. Run <code>npm install</code> to install dependencies.", body_style))
    story.append(Paragraph("3. Execute <code>npm run dev</code> for local development server.", body_style))
    story.append(Paragraph("4. Execute <code>npm run build</code> to produce production bundle in <code>/dist</code>.", body_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Generated PDF successfully: {PDF_PATH}")

if __name__ == "__main__":
    build_docx()
    build_pdf()
    print("All documents generated successfully!")
